"""
Random Forest - Clasificación de Tensión de Liquidez (modelo básico)
Ejecutar:  streamlit run app_rf_tension_liquidez.py

Configuración acordada:
  1. Excluir: ID_Observacion, Fecha, Prob_Tension_Liquidez, Brecha_Caja_90d_MXN (evitar fuga)
  2. Predictores: 21 variables de la base tal cual; Sector con OneHotEncoder
  3. Partición 70/30 aleatoria estratificada, random_state=42, sin CV
  4. class_weight='balanced'
  5. Hiperparámetros de la guía (300 árboles, max_depth=8) ajustables con sliders
  6. Métricas: accuracy, precision, recall, F1, ROC-AUC; umbral ajustable (0.50 por defecto)
  7. Salidas: exploración, métricas + matriz de confusión, curva ROC, importancia de variables
  8. Carga de Excel desde la app + predicción en lote con descarga a Excel
  9. Pestaña Reporte: descarga de un Excel con lo revisado e interpretación en lenguaje
     sencillo (Resumen, Métricas, Matriz, Umbrales, Importancia, Predicciones).
     Lectura de predicciones: alerta clara (>= umbral + 0.10), cerca del umbral
     (umbral ± 0.10) y sin alerta (< umbral - 0.10).
"""

from datetime import datetime
from io import BytesIO

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import streamlit as st
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import (
    ConfusionMatrixDisplay,
    accuracy_score,
    classification_report,
    confusion_matrix,
    f1_score,
    precision_score,
    recall_score,
    roc_auc_score,
    roc_curve,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder

# ------------------------------------------------------------------
# Parámetros fijos del modelo
# ------------------------------------------------------------------
HOJA = "Datos_Modelo"
TARGET = "Tension_Liquidez_bin"
EXCLUIR = ["ID_Observacion", "Fecha", "Prob_Tension_Liquidez", "Brecha_Caja_90d_MXN"]
CATEGORICAS = ["Sector"]
TEST_SIZE = 0.30
RANDOM_STATE = 42
ZONA_GRIS = 0.10  # ancho de "cerca del umbral" en la lectura de predicciones

DESCRIPCIONES = {
    "Sector": "Sector económico de la empresa.",
    "Ventas_12M_MXN": "Ventas de los últimos 12 meses (MXN).",
    "Crecimiento_Ventas_pct": "Variación anual de ventas.",
    "Margen_EBITDA_pct": "EBITDA / Ventas.",
    "DSO_dias": "Días promedio que tarda en cobrar a clientes.",
    "DIO_dias": "Días promedio que el inventario permanece en almacén.",
    "DPO_dias": "Días promedio que tarda en pagar a proveedores.",
    "CCC_dias": "Ciclo de conversión de efectivo: días que tarda el dinero en regresar a caja (DSO + DIO - DPO).",
    "CxC_MXN": "Cuentas por cobrar (MXN).",
    "Inventarios_MXN": "Inventarios (MXN).",
    "CxP_MXN": "Cuentas por pagar (MXN).",
    "NWC_MXN": "Capital de trabajo neto operativo (CxC + Inventarios - CxP).",
    "Deuda_CP_MXN": "Deuda financiera de corto plazo (MXN).",
    "Caja_MXN": "Efectivo disponible (MXN).",
    "Linea_Credito_Disponible_MXN": "Línea de crédito no dispuesta (MXN).",
    "Concentracion_Top5_Clientes_pct": "Participación de los 5 principales clientes en ventas.",
    "Morosidad_CxC_pct": "Proporción de cuentas por cobrar vencidas.",
    "Inventario_Obsoleto_pct": "Proporción de inventario obsoleto.",
    "Volatilidad_Ventas_pct": "Qué tanto varían las ventas.",
    "Estacionalidad_bin": "1 si la observación cae en temporada alta.",
    "EBITDA_MXN": "EBITDA estimado (MXN).",
}

st.set_page_config(page_title="RF Tensión de Liquidez", layout="wide")
st.title("Random Forest · Tensión de liquidez")
st.caption("Clasificación binaria: 1 = tensión de liquidez, 0 = situación normal. Modelo básico.")


# ------------------------------------------------------------------
# Funciones
# ------------------------------------------------------------------
@st.cache_data
def leer_excel(contenido: bytes, hoja: str | None = HOJA) -> pd.DataFrame:
    xls = pd.ExcelFile(BytesIO(contenido))
    if hoja and hoja in xls.sheet_names:
        return xls.parse(hoja)
    return xls.parse(xls.sheet_names[0])


def columnas_predictoras(df: pd.DataFrame) -> list[str]:
    return [c for c in df.columns if c not in EXCLUIR + [TARGET]]


@st.cache_resource
def entrenar(df: pd.DataFrame, n_estimators: int, max_depth: int, min_samples_leaf: int):
    predictores = columnas_predictoras(df)
    X = df[predictores]
    y = df[TARGET].astype(int)
    cat = [c for c in CATEGORICAS if c in X.columns]

    prep = ColumnTransformer(
        [("cat", OneHotEncoder(handle_unknown="ignore"), cat)],
        remainder="passthrough",
        verbose_feature_names_out=False,
    )
    modelo = RandomForestClassifier(
        n_estimators=n_estimators,
        max_depth=max_depth,
        min_samples_leaf=min_samples_leaf,
        class_weight="balanced",
        random_state=RANDOM_STATE,
        n_jobs=-1,
    )
    pipe = Pipeline([("prep", prep), ("model", modelo)])

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=TEST_SIZE, random_state=RANDOM_STATE, stratify=y
    )
    pipe.fit(X_train, y_train)
    proba_test = pipe.predict_proba(X_test)[:, 1]
    proba_train = pipe.predict_proba(X_train)[:, 1]
    return pipe, predictores, X_train, X_test, y_train, y_test, proba_train, proba_test


def metricas(y_true, proba, umbral):
    pred = (proba >= umbral).astype(int)
    return {
        "Accuracy": accuracy_score(y_true, pred),
        "Precision": precision_score(y_true, pred, zero_division=0),
        "Recall": recall_score(y_true, pred, zero_division=0),
        "F1": f1_score(y_true, pred, zero_division=0),
        "ROC-AUC": roc_auc_score(y_true, proba),
    }, pred


def a_excel(df: pd.DataFrame, hoja: str) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=hoja, index=False)
    return buffer.getvalue()


# ------------------------------------------------------------------
# Reporte descargable
# ------------------------------------------------------------------
def lectura_prediccion(p: float, umbral: float) -> str:
    if p >= umbral + ZONA_GRIS:
        return "Alerta clara"
    if p >= umbral - ZONA_GRIS:
        return "Cerca del umbral: revisar"
    return "Sin alerta"


def nivel_auc(auc: float) -> str:
    if auc >= 0.9:
        return "muy bueno"
    if auc >= 0.8:
        return "bueno"
    if auc >= 0.7:
        return "aceptable"
    if auc >= 0.6:
        return "débil"
    return "similar al azar"


def nivel_brecha(dif: float) -> str:
    if dif > 0.10:
        return "Brecha alta: el modelo rinde bastante mejor con datos conocidos (posible sobreajuste)."
    if dif > 0.05:
        return "Brecha moderada: normal en Random Forest; vigilar si crece."
    return "Brecha baja: el modelo generaliza bien."


def importancia_por_variable(pipe, predictores) -> pd.DataFrame:
    """Importancia agregada a las 21 variables originales (suma las columnas de Sector)."""
    nombres = pipe.named_steps["prep"].get_feature_names_out()
    imp = pd.Series(pipe.named_steps["model"].feature_importances_, index=nombres)
    agregada = {}
    for var in predictores:
        if var in CATEGORICAS:
            agregada[var] = imp[[n for n in nombres if n.startswith(f"{var}_")]].sum()
        else:
            agregada[var] = imp.get(var, 0.0)
    out = (pd.Series(agregada).sort_values(ascending=False).rename("Importancia")
           .reset_index().rename(columns={"index": "Variable"}))
    return out


def construir_reporte(*, nombre_archivo, n_obs, params, umbral, y_train, y_test,
                      proba_train, proba_test, m_train, m_test, pipe, predictores,
                      resultado) -> bytes:
    pred_test = (proba_test >= umbral).astype(int)
    tn, fp, fn, tp = confusion_matrix(y_test, pred_test, labels=[0, 1]).ravel()
    reales_1, reales_0 = int(tp + fn), int(tn + fp)
    alertas = int(tp + fp)
    imp = importancia_por_variable(pipe, predictores)
    top = imp.iloc[0]

    # --- Resumen
    conclusiones = [
        f"Capacidad para ordenar el riesgo: ROC-AUC de {m_test['ROC-AUC']:.3f}, nivel {nivel_auc(m_test['ROC-AUC'])} (0.5 = azar, 1.0 = perfecto).",
        f"Detecta {tp} de {reales_1} empresas con tensión real (recall {m_test['Recall']:.0%}).",
        f"Se le escapan {fn} empresas con tensión (tensión no detectada), el error más costoso en crédito.",
        f"Da {fp} falsas alarmas de {reales_0} empresas normales; de cada 100 alertas, {m_test['Precision']*100:.0f} son correctas.",
        f"La variable que más pesa es {top['Variable']} ({top['Importancia']:.0%} del total).",
        nivel_brecha(m_train["ROC-AUC"] - m_test["ROC-AUC"]).replace("Brecha", "Brecha en ROC-AUC entre entrenamiento y prueba", 1),
    ]
    resumen = pd.DataFrame({"Concepto": [f"Conclusión {i}" for i in range(1, len(conclusiones) + 1)],
                            "Valor": ["" for _ in conclusiones], "Interpretación": conclusiones})
    config = pd.DataFrame([
        ["Fecha del reporte", datetime.now().strftime("%Y-%m-%d %H:%M"), "Momento en que se generó el reporte."],
        ["Archivo de entrenamiento", nombre_archivo, "Base usada para entrenar y evaluar."],
        ["Observaciones", n_obs, "Número de empresas en la base."],
        ["Entrenamiento / prueba", f"{len(y_train)} / {len(y_test)}", "70% para aprender y 30% para calificar (empresas que el modelo no vio)."],
        ["Número de árboles", params["n_estimators"], "Cuántos árboles (analistas) forman el comité."],
        ["Profundidad máxima", params["max_depth"], "Máximo de preguntas seguidas que hace cada árbol."],
        ["Mínimo de obs. por hoja", params["min_samples_leaf"], "Empresas mínimas para que un árbol saque una conclusión."],
        ["Umbral de decisión", umbral, "Probabilidad a partir de la cual se clasifica como Tensión."],
        ["Balance de clases", "class_weight='balanced'", "Da el mismo peso a ambas clases aunque una sea más frecuente."],
        ["Variables excluidas", ", ".join(EXCLUIR), "Se excluyen para evitar fuga de información."],
    ], columns=["Concepto", "Valor", "Interpretación"])

    # --- Métricas
    preguntas = {
        "Accuracy": ("De todas las empresas, ¿cuántas clasificó bien?",
                     f"Clasificó bien {tp + tn} de {len(y_test)} empresas."),
        "Precision": ("Cuando dice Tensión, ¿qué tan seguido acierta?",
                      f"De {alertas} alertas, {tp} fueron correctas."),
        "Recall": ("De las empresas con tensión real, ¿cuántas detectó?",
                   f"Detectó {tp} de {reales_1}."),
        "F1": ("Balance entre precision y recall.", "Una sola nota que combina ambas."),
        "ROC-AUC": ("¿Qué tan bien ordena de más a menos riesgosa?",
                    f"Nivel {nivel_auc(m_test['ROC-AUC'])}. No depende del umbral."),
    }
    metricas_df = pd.DataFrame([
        [k, round(m_test[k], 4), round(m_train[k], 4), round(m_train[k] - m_test[k], 4),
         preguntas[k][0], preguntas[k][1] + " " + nivel_brecha(m_train[k] - m_test[k])]
        for k in m_test
    ], columns=["Métrica", "Prueba", "Entrenamiento", "Diferencia", "Pregunta que responde", "Interpretación"])

    # --- Matriz
    matriz_df = pd.DataFrame([
        ["Real Normal / Predicho Normal", "Acierto", tn, tn / len(y_test), "Empresas normales identificadas correctamente."],
        ["Real Normal / Predicho Tensión", "Falsa alarma (falso positivo)", fp, fp / len(y_test), "Era normal y el modelo dijo tensión. Costo: una revisión de más."],
        ["Real Tensión / Predicho Normal", "Tensión no detectada (falso negativo)", fn, fn / len(y_test), "Tenía tensión y el modelo dijo normal. Es el error más caro en crédito."],
        ["Real Tensión / Predicho Tensión", "Acierto", tp, tp / len(y_test), "Empresas con tensión detectadas correctamente."],
    ], columns=["Cuadro", "Tipo", "Empresas", "% del total de prueba", "Interpretación"])

    # --- Umbrales
    filas = []
    for u in np.round(np.arange(0.10, 0.901, 0.05), 2):
        m_u, p_u = metricas(y_test, proba_test, u)
        tn_u, fp_u, fn_u, tp_u = confusion_matrix(y_test, p_u, labels=[0, 1]).ravel()
        filas.append([u, tp_u, fn_u, fp_u, round(m_u["Accuracy"], 4), round(m_u["Precision"], 4),
                      round(m_u["Recall"], 4), round(m_u["F1"], 4),
                      f"Detecta {tp_u} de {reales_1}; se escapan {fn_u}; {fp_u} falsas alarmas."
                      + ("  ← UMBRAL VIGENTE" if abs(u - umbral) < 1e-9 else "")])
    umbrales_df = pd.DataFrame(filas, columns=["Umbral", "Tensión detectada", "Tensión no detectada",
                                               "Falsas alarmas", "Accuracy", "Precision", "Recall", "F1",
                                               "Interpretación"])
    if not np.isclose(umbrales_df["Umbral"], umbral).any():
        m_u, p_u = metricas(y_test, proba_test, umbral)
        tn_u, fp_u, fn_u, tp_u = confusion_matrix(y_test, p_u, labels=[0, 1]).ravel()
        umbrales_df.loc[len(umbrales_df)] = [umbral, tp_u, fn_u, fp_u, round(m_u["Accuracy"], 4),
                                             round(m_u["Precision"], 4), round(m_u["Recall"], 4),
                                             round(m_u["F1"], 4),
                                             f"Detecta {tp_u} de {reales_1}; se escapan {fn_u}; {fp_u} falsas alarmas.  ← UMBRAL VIGENTE"]
        umbrales_df = umbrales_df.sort_values("Umbral").reset_index(drop=True)

    # --- Importancia
    imp["Importancia_acumulada"] = imp["Importancia"].cumsum()
    imp.insert(0, "Rank", range(1, len(imp) + 1))
    imp["Descripción"] = imp["Variable"].map(DESCRIPCIONES).fillna("")
    imp["Interpretación"] = [
        "Variable más usada por el modelo." if i == 0 else
        ("Entre las variables que explican el 80% de la decisión." if acum - w < 0.80 else "Peso menor en la decisión.")
        for i, (w, acum) in enumerate(zip(imp["Importancia"], imp["Importancia_acumulada"]))
    ]

    # --- Escritura con formato
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as w:
        pd.DataFrame().to_excel(w, sheet_name="Resumen")
        config.to_excel(w, sheet_name="Resumen", index=False, startrow=2)
        resumen.to_excel(w, sheet_name="Resumen", index=False, startrow=len(config) + 6)
        metricas_df.to_excel(w, sheet_name="Métricas", index=False, startrow=2)
        matriz_df.to_excel(w, sheet_name="Matriz", index=False, startrow=2)
        umbrales_df.to_excel(w, sheet_name="Umbrales", index=False, startrow=2)
        imp.to_excel(w, sheet_name="Importancia", index=False, startrow=2)
        if resultado is not None:
            pred = resultado.copy()
            pred["Lectura"] = [lectura_prediccion(p, umbral) for p in pred["Prob_Tension_Modelo"]]
            clave = ["Prob_Tension_Modelo", "Clasificacion", "Lectura", "Tension_Predicha", "Umbral_usado"]
            ids = [c for c in pred.columns if c not in clave and c not in predictores]
            pred = pred[ids + ["Sector"] + clave + [c for c in predictores if c != "Sector"]]
            pred.to_excel(w, sheet_name="Predicciones", index=False, startrow=2)
        _formatear(w.book, umbral, len(config))
    return buffer.getvalue()


def _formatear(wb, umbral, n_config):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    titulos = {
        "Resumen": "Reporte · Random Forest de tensión de liquidez · Configuración",
        "Métricas": "Métricas del modelo (prueba = empresas no vistas al entrenar)",
        "Matriz": f"Matriz de confusión (umbral {umbral:.2f})",
        "Umbrales": "Efecto de mover el umbral de decisión (conjunto de prueba)",
        "Importancia": "Importancia de las 21 variables (suma 100%)",
        "Predicciones": f"Predicciones de empresas nuevas (umbral {umbral:.2f}; zona 'cerca del umbral' ± {ZONA_GRIS:.2f})",
    }
    head_fill = PatternFill("solid", fgColor="14213D")
    vigente_fill = PatternFill("solid", fgColor="FBE3D3")
    lectura_fill = {"Alerta clara": "F8CBAD", "Cerca del umbral: revisar": "FFF2CC", "Sin alerta": "E2EFDA"}
    pct_cols = {"% del total de prueba", "Importancia", "Importancia_acumulada"}

    for ws in wb.worksheets:
        ws["A1"] = titulos.get(ws.title, ws.title)
        ws["A1"].font = Font(bold=True, size=14, color="14213D")
        header_rows = [3] + ([n_config + 7] if ws.title == "Resumen" else [])
        if ws.title == "Resumen":
            ws.cell(row=n_config + 6, column=1, value="Conclusiones clave").font = Font(bold=True, size=14, color="14213D")
        for r in header_rows:
            for c in ws[r]:
                if c.value is not None:
                    c.font = Font(bold=True, color="FFFFFF")
                    c.fill = head_fill
                    c.alignment = Alignment(wrap_text=True, vertical="center")
        headers = {c.column: c.value for c in ws[3]}
        for col_idx, name in headers.items():
            letra = get_column_letter(col_idx)
            largo = max(len(str(name or "")), *(len(str(c.value)) for c in ws[letra][3:] if c.value is not None), 8)
            ancho = 70 if name in ("Interpretación", "Pregunta que responde", "Descripción") else min(largo + 2, 40)
            ws.column_dimensions[letra].width = ancho
            for c in ws[letra][3:]:
                c.alignment = Alignment(wrap_text=True, vertical="top",
                                        horizontal="left" if name == "Valor" else None)
                if name in pct_cols and isinstance(c.value, (int, float)):
                    c.number_format = "0.0%"
                if name == "Prob_Tension_Modelo" and isinstance(c.value, (int, float)):
                    c.number_format = "0.0%"
        if ws.title == "Umbrales":
            for row in ws.iter_rows(min_row=4):
                if isinstance(row[0].value, (int, float)) and abs(row[0].value - umbral) < 1e-9:
                    for c in row:
                        c.fill = vigente_fill
                        c.font = Font(bold=True)
        if ws.title == "Predicciones":
            col_lect = [k for k, v in headers.items() if v == "Lectura"]
            if col_lect:
                for c in ws[get_column_letter(col_lect[0])][3:]:
                    if c.value in lectura_fill:
                        c.fill = PatternFill("solid", fgColor=lectura_fill[c.value])
        ws.freeze_panes = "A4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True


# ------------------------------------------------------------------
# Barra lateral: datos e hiperparámetros
# ------------------------------------------------------------------
st.sidebar.header("1. Datos")
archivo = st.sidebar.file_uploader("Base de entrenamiento (.xlsx)", type=["xlsx", "xlsm"])

st.sidebar.header("2. Hiperparámetros")
n_estimators = st.sidebar.slider("Número de árboles (n_estimators)", 50, 1000, 300, 50)
max_depth = st.sidebar.slider("Profundidad máxima (max_depth)", 2, 20, 8, 1)
min_samples_leaf = st.sidebar.slider("Mínimo de obs. por hoja (min_samples_leaf)", 1, 20, 1, 1)

st.sidebar.header("3. Umbral de decisión")
umbral = st.sidebar.slider("Clasificar como tensión si probabilidad ≥", 0.05, 0.95, 0.50, 0.01)

st.sidebar.markdown("---")
st.sidebar.caption(
    f"Partición {int((1-TEST_SIZE)*100)}/{int(TEST_SIZE*100)} estratificada · "
    f"random_state={RANDOM_STATE} · class_weight='balanced'"
)

if archivo is None:
    st.info("Sube el archivo Excel con la hoja **Datos_Modelo** para entrenar el modelo.")
    st.stop()

df = leer_excel(archivo.getvalue())

faltantes = [c for c in [TARGET] + CATEGORICAS if c not in df.columns]
if faltantes:
    st.error(f"No se encontraron las columnas requeridas: {', '.join(faltantes)}")
    st.stop()

predictores = columnas_predictoras(df)
if df[predictores + [TARGET]].isna().any().any():
    st.warning("La base tiene valores faltantes; se eliminan esas filas para entrenar.")
    df = df.dropna(subset=predictores + [TARGET])

with st.spinner("Entrenando Random Forest..."):
    pipe, predictores, X_train, X_test, y_train, y_test, proba_train, proba_test = entrenar(
        df, n_estimators, max_depth, min_samples_leaf
    )

m_test, pred_test = metricas(y_test, proba_test, umbral)
m_train, _ = metricas(y_train, proba_train, umbral)

tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs(
    ["Exploración", "Desempeño", "Curva ROC", "Importancia de variables", "Predicción nuevas empresas",
     "Reporte"]
)
resultado = None  # se llena si se suben empresas nuevas en la pestaña 5

# ------------------------------------------------------------------
# 1. Exploración
# ------------------------------------------------------------------
with tab1:
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Observaciones", f"{len(df):,}")
    c2.metric("Predictores", len(predictores))
    c3.metric("% con tensión (1)", f"{df[TARGET].mean():.1%}")
    c4.metric("Sectores", df["Sector"].nunique())

    st.markdown(f"**Variables excluidas:** {', '.join([c for c in EXCLUIR if c in df.columns])}")
    st.subheader("Vista previa")
    st.dataframe(df.head(20))

    g1, g2 = st.columns(2)
    with g1:
        st.subheader("Distribución del target")
        conteo = df[TARGET].value_counts().sort_index()
        fig, ax = plt.subplots(figsize=(5, 3.5))
        ax.bar(["0 · Normal", "1 · Tensión"], conteo.values, color=["#4C78A8", "#E45756"])
        for i, v in enumerate(conteo.values):
            ax.text(i, v, f"{v} ({v/conteo.sum():.0%})", ha="center", va="bottom")
        ax.set_ylabel("Observaciones")
        ax.spines[["top", "right"]].set_visible(False)
        st.pyplot(fig)
    with g2:
        st.subheader("Tasa de tensión por sector")
        por_sector = (
            df.groupby("Sector")[TARGET].agg(Observaciones="count", Tasa_tension="mean")
            .sort_values("Tasa_tension")
        )
        fig, ax = plt.subplots(figsize=(5, 3.5))
        ax.barh(por_sector.index, por_sector["Tasa_tension"], color="#E45756")
        for i, v in enumerate(por_sector["Tasa_tension"]):
            ax.text(v, i, f" {v:.0%}", va="center")
        ax.set_xlim(0, 1)
        ax.set_xlabel("% con tensión de liquidez")
        ax.spines[["top", "right"]].set_visible(False)
        st.pyplot(fig)

    st.subheader("Estadísticos descriptivos de los predictores numéricos")
    st.dataframe(df[predictores].describe().T.round(3))

# ------------------------------------------------------------------
# 2. Desempeño
# ------------------------------------------------------------------
with tab2:
    st.subheader(f"Métricas en el conjunto de prueba (umbral = {umbral:.2f})")
    cols = st.columns(5)
    for col, (nombre, valor) in zip(cols, m_test.items()):
        col.metric(nombre, f"{valor:.3f}", delta=f"{valor - m_train[nombre]:+.3f} vs entrenamiento",
                   delta_color="off")
    st.caption(
        f"Entrenamiento: {len(y_train)} obs. · Prueba: {len(y_test)} obs. "
        "Una diferencia grande entre entrenamiento y prueba indica sobreajuste."
    )

    g1, g2 = st.columns(2)
    with g1:
        st.subheader("Matriz de confusión")
        cm = confusion_matrix(y_test, pred_test, labels=[0, 1])
        fig, ax = plt.subplots(figsize=(4.5, 4))
        ConfusionMatrixDisplay(cm, display_labels=["0 · Normal", "1 · Tensión"]).plot(
            ax=ax, cmap="Blues", colorbar=False, values_format="d"
        )
        ax.set_xlabel("Predicción")
        ax.set_ylabel("Real")
        st.pyplot(fig)
        tn, fp, fn, tp = cm.ravel()
        st.markdown(
            f"- **Falsos negativos (tensión no detectada):** {fn}\n"
            f"- **Falsos positivos (falsa alarma):** {fp}"
        )
    with g2:
        st.subheader("Reporte de clasificación")
        reporte = pd.DataFrame(
            classification_report(
                y_test, pred_test, target_names=["0 · Normal", "1 · Tensión"],
                output_dict=True, zero_division=0,
            )
        ).T.round(3)
        st.dataframe(reporte)

# ------------------------------------------------------------------
# 3. Curva ROC
# ------------------------------------------------------------------
with tab3:
    st.subheader(f"Curva ROC · AUC prueba = {m_test['ROC-AUC']:.3f}")
    fpr, tpr, thr = roc_curve(y_test, proba_test)
    fpr_tr, tpr_tr, _ = roc_curve(y_train, proba_train)
    idx = np.argmin(np.abs(thr - umbral))

    fig, ax = plt.subplots(figsize=(6, 5))
    ax.plot(fpr, tpr, color="#E45756", lw=2, label=f"Prueba (AUC {m_test['ROC-AUC']:.3f})")
    ax.plot(fpr_tr, tpr_tr, color="#4C78A8", lw=1, ls="--",
            label=f"Entrenamiento (AUC {m_train['ROC-AUC']:.3f})")
    ax.plot([0, 1], [0, 1], color="grey", lw=1, ls=":", label="Azar")
    ax.scatter(fpr[idx], tpr[idx], color="black", zorder=5, label=f"Umbral {umbral:.2f}")
    ax.set_xlabel("Tasa de falsos positivos")
    ax.set_ylabel("Tasa de verdaderos positivos (recall)")
    ax.legend(loc="lower right")
    ax.spines[["top", "right"]].set_visible(False)
    st.pyplot(fig)

# ------------------------------------------------------------------
# 4. Importancia de variables
# ------------------------------------------------------------------
with tab4:
    st.subheader("Importancia de variables (feature_importances_)")
    nombres = pipe.named_steps["prep"].get_feature_names_out()
    imp = pd.DataFrame(
        {"Variable": nombres, "Importancia": pipe.named_steps["model"].feature_importances_}
    ).sort_values("Importancia", ascending=False).reset_index(drop=True)

    top_n = st.slider("Número de variables a mostrar", 5, len(imp), min(15, len(imp)))
    top = imp.head(top_n).iloc[::-1]
    fig, ax = plt.subplots(figsize=(7, 0.35 * top_n + 1))
    ax.barh(top["Variable"], top["Importancia"], color="#4C78A8")
    ax.set_xlabel("Importancia (reducción media de impureza)")
    ax.spines[["top", "right"]].set_visible(False)
    st.pyplot(fig)
    st.caption(
        "Nota: CCC_dias es combinación exacta de DSO, DIO y DPO, y NWC_MXN de CxC, Inventarios y CxP; "
        "la importancia se reparte entre variables redundantes."
    )
    st.dataframe(imp.assign(Importancia=imp["Importancia"].round(4)))

# ------------------------------------------------------------------
# 5. Predicción en lote
# ------------------------------------------------------------------
with tab5:
    st.subheader("Predicción para empresas nuevas")
    st.markdown(
        f"Sube un Excel con las mismas columnas predictoras ({len(predictores)}). "
        "Puede incluir columnas adicionales (ID, nombre, etc.); se conservan en el resultado."
    )
    plantilla = pd.DataFrame(columns=["ID_Empresa"] + predictores)
    st.download_button(
        "Descargar plantilla (.xlsx)",
        data=a_excel(plantilla, "Nuevas_Empresas"),
        file_name="plantilla_nuevas_empresas.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    nuevo = st.file_uploader("Archivo de empresas nuevas (.xlsx)", type=["xlsx", "xlsm"], key="nuevo")
    if nuevo is not None:
        df_nuevo = leer_excel(nuevo.getvalue(), hoja=None)
        faltan = [c for c in predictores if c not in df_nuevo.columns]
        if faltan:
            st.error(f"Faltan columnas en el archivo: {', '.join(faltan)}")
        elif df_nuevo[predictores].isna().any().any():
            st.error("El archivo tiene valores vacíos en columnas predictoras; complétalos antes de predecir.")
        else:
            sectores_nuevos = set(df_nuevo["Sector"]) - set(df["Sector"])
            if sectores_nuevos:
                st.warning(f"Sectores no vistos en entrenamiento (se tratan como 'ninguno'): {sectores_nuevos}")
            proba = pipe.predict_proba(df_nuevo[predictores])[:, 1]
            resultado = df_nuevo.copy()
            resultado["Prob_Tension_Modelo"] = proba.round(4)
            resultado["Tension_Predicha"] = (proba >= umbral).astype(int)
            resultado["Clasificacion"] = np.where(resultado["Tension_Predicha"] == 1, "Tensión", "Normal")
            resultado["Umbral_usado"] = umbral

            c1, c2 = st.columns(2)
            c1.metric("Empresas evaluadas", len(resultado))
            c2.metric("Con tensión predicha", f"{resultado['Tension_Predicha'].sum()} "
                      f"({resultado['Tension_Predicha'].mean():.0%})")
            st.dataframe(resultado)
            st.download_button(
                "Descargar predicciones (.xlsx)",
                data=a_excel(resultado, "Predicciones"),
                file_name="predicciones_tension_liquidez.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

# ------------------------------------------------------------------
# 6. Reporte descargable
# ------------------------------------------------------------------
with tab6:
    st.subheader("Reporte de lo revisado (.xlsx)")
    st.markdown(
        "Genera un Excel con la configuración vigente, los resultados y una columna de "
        "**interpretación en lenguaje sencillo**. Refleja los sliders y el umbral actuales."
    )
    hojas = ["Resumen", "Métricas", "Matriz", "Umbrales", "Importancia"]
    if resultado is not None:
        hojas.append("Predicciones")
    st.markdown("**Hojas incluidas:** " + " · ".join(hojas))
    if resultado is None:
        st.caption("Para incluir la hoja Predicciones, sube primero el archivo de empresas nuevas en la pestaña anterior.")
    st.caption(
        f"Lectura de predicciones: Alerta clara (≥ {min(umbral + ZONA_GRIS, 1):.2f}) · "
        f"Cerca del umbral: revisar ({max(umbral - ZONA_GRIS, 0):.2f} a {min(umbral + ZONA_GRIS, 1):.2f}) · "
        f"Sin alerta (< {max(umbral - ZONA_GRIS, 0):.2f})."
    )
    reporte_xlsx = construir_reporte(
        nombre_archivo=archivo.name if hasattr(archivo, "name") else "archivo cargado",
        n_obs=len(df),
        params={"n_estimators": n_estimators, "max_depth": max_depth, "min_samples_leaf": min_samples_leaf},
        umbral=umbral, y_train=y_train, y_test=y_test, proba_train=proba_train, proba_test=proba_test,
        m_train=m_train, m_test=m_test, pipe=pipe, predictores=predictores, resultado=resultado,
    )
    st.download_button(
        "Descargar reporte (.xlsx)",
        data=reporte_xlsx,
        file_name=f"reporte_rf_tension_liquidez_{datetime.now():%Y%m%d_%H%M}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
    )
